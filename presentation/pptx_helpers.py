"""Hjälpfunktioner för att bygga kursens presentation med python-pptx.

Presentationen genereras av bygg_presentation.py. Ändra innehållet där och kör:

    pip install python-pptx
    python presentation/bygg_presentation.py
"""
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# Färger
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
TEAL_LIGHT = RGBColor(0xD9, 0xEE, 0xEF)
ORANGE = RGBColor(0xE8, 0x77, 0x2E)
ORANGE_LIGHT = RGBColor(0xFC, 0xEB, 0xDD)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
GREY = RGBColor(0x5B, 0x65, 0x77)
GREY_LIGHT = RGBColor(0xD5, 0xDA, 0xE1)
DARK = RGBColor(0x22, 0x28, 0x31)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
GREEN_LIGHT = RGBColor(0xDF, 0xF2, 0xE6)
RED = RGBColor(0xC0, 0x39, 0x2B)
RED_LIGHT = RGBColor(0xF8, 0xE0, 0xDD)
CODE_BG = RGBColor(0xF4, 0xF6, 0xF8)
CODE_COMMENT = RGBColor(0x6A, 0x73, 0x7D)
CODE_KEYWORD = RGBColor(0x0E, 0x6B, 0x74)
CODE_STRING = RGBColor(0xA6, 0x4B, 0x00)

FONT = "Calibri"
MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_TOP = Inches(1.45)
CONTENT_W = SLIDE_W - 2 * MARGIN

FSH_KEYWORDS = (
    "Alias:", "Profile:", "Parent:", "Id:", "Title:", "Description:", "Extension:", "Context:",
    "ValueSet:", "CodeSystem:", "Instance:", "InstanceOf:", "Usage:", "Invariant:", "Expression:",
    "Severity:", "RuleSet:", "Logical:", "Mapping:", "Source:", "Target:",
)


class Deck:
    def __init__(self, footer_title):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.footer_title = footer_title
        self.section = ""
        self.count = 0

    def save(self, path):
        self.prs.save(path)

    # ------------------------------------------------------------------ grundläggande
    def _slide(self, notes=None, background=None):
        slide = self.prs.slides.add_slide(self.blank)
        self.count += 1
        if background is not None:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = background
        if notes:
            slide.notes_slide.notes_text_frame.text = notes.strip()
        return slide

    def _footer(self, slide, dark=False):
        color = GREY_LIGHT if dark else GREY
        text = self.footer_title + (f"  ·  {self.section}" if self.section else "")
        tb = slide.shapes.add_textbox(MARGIN, Inches(7.02), Inches(10), Inches(0.35))
        _set_text(tb.text_frame, text, size=10, color=color)
        nb = slide.shapes.add_textbox(SLIDE_W - MARGIN - Inches(1), Inches(7.02), Inches(1), Inches(0.35))
        _set_text(nb.text_frame, str(self.count), size=10, color=color, align=PP_ALIGN.RIGHT)

    def content_slide(self, title, notes=None, subtitle=None):
        slide = self._slide(notes)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
        _fill(bar, TEAL)
        tb = slide.shapes.add_textbox(MARGIN, Inches(0.35), CONTENT_W, Inches(0.75))
        _set_text(tb.text_frame, title, size=30, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)
        if subtitle:
            st = slide.shapes.add_textbox(MARGIN, Inches(0.98), CONTENT_W, Inches(0.4))
            _set_text(st.text_frame, subtitle, size=16, color=TEAL)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, MARGIN, Inches(1.36), SLIDE_W - MARGIN, Inches(1.36))
        line.line.color.rgb = GREY_LIGHT
        line.line.width = Pt(1)
        self._footer(slide)
        return slide

    def title_slide(self, title, subtitle, details, notes=None):
        slide = self._slide(notes, background=NAVY)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.2), SLIDE_W, Inches(0.12))
        _fill(band, ORANGE)
        tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.8))
        _set_text(tb.text_frame, title, size=44, bold=True, color=WHITE)
        st = slide.shapes.add_textbox(Inches(0.9), Inches(3.4), Inches(11.5), Inches(1.2))
        _set_text(st.text_frame, subtitle, size=22, color=TEAL_LIGHT)
        dt = slide.shapes.add_textbox(Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.0))
        _set_text(dt.text_frame, details, size=16, color=GREY_LIGHT)
        return slide

    def module_slide(self, number, title, goals, agenda, notes=None):
        """Avdelare för en modul: nummer, titel, mål och tidsplan."""
        self.section = f"Modul {number} – {title}"
        slide = self._slide(notes, background=NAVY)
        num = slide.shapes.add_textbox(Inches(0.9), Inches(0.7), Inches(3), Inches(1.6))
        _set_text(num.text_frame, f"{number}", size=96, bold=True, color=ORANGE)
        tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.25), Inches(11.5), Inches(1.0))
        _set_text(tb.text_frame, title, size=36, bold=True, color=WHITE)

        gl = slide.shapes.add_textbox(Inches(0.9), Inches(3.45), Inches(7.4), Inches(3.3))
        tf = gl.text_frame
        tf.word_wrap = True
        _set_text(tf, "Efter modulen kan du …", size=16, bold=True, color=TEAL_LIGHT)
        for goal in goals:
            p = tf.add_paragraph()
            _run(p, "›  " + goal, size=16, color=WHITE)
            p.space_before = Pt(6)

        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(3.45), Inches(3.9), Inches(2.9))
        _fill(box, RGBColor(0x26, 0x3A, 0x62))
        box.line.fill.background()
        box.adjustments[0] = 0.06
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        _set_text(tf, "50 minuter", size=16, bold=True, color=ORANGE)
        for minutes, label in agenda:
            p = tf.add_paragraph()
            p.space_before = Pt(8)
            _run(p, f"{minutes:>2} min  ", size=15, bold=True, color=WHITE)
            _run(p, label, size=15, color=GREY_LIGHT)
        self._footer(slide, dark=True)
        return slide

    def exercise_slide(self, title, items, where=None, notes=None, minutes=None):
        """Övningsbild med orange banner."""
        slide = self._slide(notes)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.2))
        _fill(band, ORANGE)
        label = "ÖVNING" + (f"  ·  {minutes} min" if minutes else "")
        lb = slide.shapes.add_textbox(MARGIN, Inches(0.12), Inches(6), Inches(0.4))
        _set_text(lb.text_frame, label, size=14, bold=True, color=WHITE)
        tb = slide.shapes.add_textbox(MARGIN, Inches(0.42), CONTENT_W, Inches(0.7))
        _set_text(tb.text_frame, title, size=28, bold=True, color=WHITE)
        top = Inches(1.55)
        if where:
            wb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, top, CONTENT_W, Inches(0.55))
            _fill(wb, ORANGE_LIGHT)
            wb.line.fill.background()
            wb.adjustments[0] = 0.2
            wtf = wb.text_frame
            wtf.margin_left = Inches(0.2)
            _set_text(wtf, "📂  " + where, size=15, color=DARK, font=FONT)
            top += Inches(0.75)
        self.bullets(slide, items, MARGIN, top, CONTENT_W, Inches(6.85) - top, size=19)
        self._footer(slide)
        return slide

    # ------------------------------------------------------------------ innehåll
    def bullets(self, slide, items, left, top, width, height, size=20, color=DARK, spacing=8):
        """Punktlista. Ett element kan vara en sträng, eller (nivå, text). **fet** stöds."""
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for item in items:
            level, text = item if isinstance(item, tuple) else (0, item)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(spacing if level == 0 else 3)
            indent = Inches(0.32 + 0.35 * level)
            _bullet(p, indent, "•" if level == 0 else "–", TEAL if level == 0 else GREY)
            _rich(p, text, size=size - 2 * level, color=color if level == 0 else GREY)
        return tb

    def text(self, slide, text, left, top, width, height, size=18, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            _rich(p, line, size=size, color=color, bold=bold, italic=italic)
        return tb

    def code(self, slide, code, left, top, width, height, size=13, lang="fsh", title=None):
        """Kodruta med enkel syntaxfärgning."""
        if title:
            cap = slide.shapes.add_textbox(left, top, width, Inches(0.3))
            _set_text(cap.text_frame, title, size=12, bold=True, color=GREY)
            top += Inches(0.32)
            height -= Inches(0.32)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        _fill(box, CODE_BG)
        box.line.color.rgb = GREY_LIGHT
        box.line.width = Pt(0.75)
        edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
        _fill(edge, TEAL)
        tf = box.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.22)
        tf.margin_top = Inches(0.1)
        tf.margin_right = Inches(0.1)
        lines = code.strip("\n").split("\n")
        # Krymp texten så att den får plats (teckenbredd ≈ 0,6 em, radhöjd ≈ 1,2 em)
        usable_w = (width - Inches(0.35)) / 914400 * 72
        usable_h = (height - Inches(0.2)) / 914400 * 72
        longest = max(len(line) for line in lines) or 1
        size = min(size, usable_w / (0.6 * longest), usable_h / (1.22 * len(lines)))
        size = max(round(size * 2) / 2, 8)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.0
            for txt, color, bold in _highlight(line, lang):
                _run(p, txt, size=size, color=color, bold=bold, font=MONO)
        return box

    def table(self, slide, rows, left, top, width, col_widths=None, size=14, height=None,
              header_fill=NAVY, first_col_bold=False, row_height=None):
        n_rows, n_cols = len(rows), len(rows[0])
        row_h = row_height or Inches(0.42)
        shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height or row_h * n_rows)
        tbl = shape.table
        tbl.first_row = True
        if col_widths:
            total = sum(col_widths)
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = Emu(int(width * w / total))
        for r, row in enumerate(rows):
            tbl.rows[r].height = row_h
            for c, value in enumerate(row):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if r == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_fill
                    _rich(p, value, size=size, color=WHITE, bold=True)
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
                    _rich(p, value, size=size, color=DARK, bold=(first_col_bold and c == 0))
        return shape

    def box(self, slide, text, left, top, width, height, fill=TEAL_LIGHT, color=DARK, size=16,
            bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE):
        shp = slide.shapes.add_shape(shape, left, top, width, height)
        _fill(shp, fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(1.25)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = 0.12
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.1)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        for i, line_text in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            _rich(p, line_text, size=size, color=color, bold=bold)
        return shp

    def arrow(self, slide, x1, y1, x2, y2, color=GREY, width=2.0):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        conn.line.color.rgb = color
        conn.line.width = Pt(width)
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
        return conn

    def callout(self, slide, text, left, top, width, height, kind="info", size=15):
        fills = {"info": (TEAL_LIGHT, TEAL), "warn": (ORANGE_LIGHT, ORANGE), "ok": (GREEN_LIGHT, GREEN),
                 "bad": (RED_LIGHT, RED)}
        fill, edge_color = fills[kind]
        shp = self.box(slide, text, left, top, width, height, fill=fill, size=size, align=PP_ALIGN.LEFT,
                       shape=MSO_SHAPE.RECTANGLE)
        shp.text_frame.margin_left = Inches(0.22)
        edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.07), height)
        _fill(edge, edge_color)
        return shp


# ---------------------------------------------------------------------- lågnivå
def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _run(p, text, size=18, color=DARK, bold=False, italic=False, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return r


def _set_text(tf, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=None, font=FONT):
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, color=color, bold=bold, font=font)


def _rich(p, text, size=18, color=DARK, bold=False, italic=False):
    """Stöd för **fetstil** och `kod` i löptext."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _run(p, part[2:-2], size=size, color=color, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            _run(p, part[1:-1], size=size - 1, color=CODE_KEYWORD, bold=bold, font=MONO)
        else:
            _run(p, part, size=size, color=color, bold=bold, italic=italic)


def _bullet(p, indent, char, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent)))
    pPr.set("indent", str(-int(Inches(0.28))))
    for tag in ("a:buNone", "a:buChar", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    clr.append(srgb)
    pPr.append(clr)
    bu = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bu)


def _highlight(line, lang):
    """Mycket enkel syntaxfärgning: kommentarer, nyckelord och strängar."""
    if lang in ("fsh", "java", "json"):
        m = re.search(r"(^|\s)//", line)
        if m and line[:m.start()].count('"') % 2 == 0:
            idx = m.start() + len(m.group(1))
            return _highlight(line[:idx], lang) + [(line[idx:], CODE_COMMENT, False)]
    if lang == "fsh":
        stripped = line.lstrip()
        for kw in FSH_KEYWORDS:
            if stripped.startswith(kw):
                pre = line[: len(line) - len(stripped)]
                return [(pre + kw, CODE_KEYWORD, True)] + _strings(stripped[len(kw):])
    if lang == "json":
        m = re.match(r'^(\s*)("[^"]+")(\s*:)(.*)$', line)
        if m:
            return [(m.group(1), DARK, False), (m.group(2), CODE_KEYWORD, False), (m.group(3), DARK, False)] + _strings(m.group(4))
    if lang == "http":
        m = re.match(r"^(GET|POST|PUT|PATCH|DELETE)(\s.*)$", line)
        if m:
            return [(m.group(1), CODE_KEYWORD, True), (m.group(2), DARK, False)]
    return _strings(line) if line else [(" ", DARK, False)]


def _strings(text):
    out = []
    for part in re.split(r'("[^"]*")', text):
        if not part:
            continue
        out.append((part, CODE_STRING if part.startswith('"') else DARK, False))
    return out or [("", DARK, False)]
